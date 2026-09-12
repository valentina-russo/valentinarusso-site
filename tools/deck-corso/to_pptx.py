#!/usr/bin/env python3
"""
deck.html -> corso-base-hd.pptx

Misura ogni slide nel browser (Playwright) e ricostruisce la pagina in
PowerPoint con python-pptx: i testi diventano caselle di testo modificabili
nella posizione esatta, con font, corpo, colore e allineamento letti dal
CSS calcolato; grafici, foto, cornici e anelli diventano immagini ritagliate
dallo screenshot dell'elemento. Il PPTX si importa in Canva senza perdere
l'impaginato.

    py tools/deck-corso/to_pptx.py            -> tools/deck-corso/out/corso-base-hd.pptx
"""
from __future__ import annotations

import io
import re
from pathlib import Path

from PIL import Image
from playwright.sync_api import sync_playwright
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Pt

QUI = Path(__file__).resolve().parent
SORGENTE = QUI / "deck.html"
USCITA = QUI / "out" / "corso-base-hd.pptx"

W, H = 1920, 1080
LARGH_IN = 13.3333
EMU_PER_PX = LARGH_IN * 914400 / W

FONT = {"Playfair": "Playfair Display", "Outfit": "Outfit", "Fell": "IM Fell English"}

# Elementi che diventano immagini (screenshot), con tutto quello che contengono
RASTER = [".telaio-chart", ".cornice", ".duotone", ".ovale", "svg.anello", "svg.anello-grande",
          "svg.anello-agenda", "svg.passi", "img.marchio"]
# Elementi di testo: foglie con testo proprio
TESTO = ("h1, h2, p, li, figcaption, span.nota-nostra, span.scadenza, span.link, span.num-slide, "
         "span.firma, .scala b, .scala span, .legenda span, .prezzi div, .prezzi b, .passo")

JS_MISURA = r"""
(sel) => {
  const out = [];
  document.querySelectorAll('section.slide').forEach((sec, si) => {
    const sr = sec.getBoundingClientRect();
    const cs = getComputedStyle(sec);
    const slide = {i: si, bg: cs.backgroundColor, nome: sec.dataset.nome, raster: [], testi: [], linee: [], punti: []};
    sec.querySelectorAll(sel.raster).forEach(el => {
      const r = el.getBoundingClientRect();
      if (r.width < 2 || r.height < 2) return;
      if (el.tagName === 'IMG' && el.closest('.telaio-chart, .cornice, .duotone, .ovale')) return;
      slide.raster.push({x: r.left - sr.left, y: r.top - sr.top, w: r.width, h: r.height});
    });
    sec.querySelectorAll(sel.testo).forEach(el => {
      if (el.closest(sel.raster)) return;
      let own = '';
      el.childNodes.forEach(n => {
        if (n.nodeType === 3) own += n.textContent;
        else if (n.nodeType === 1 && n.tagName !== 'BR' && !n.matches(sel.testo)) own += n.innerText;
        else if (n.nodeType === 1 && n.tagName === 'BR') own += '\n';
      });
      own = own.replace(/[ \t]+/g, ' ').replace(/ *\n */g, '\n').trim();
      if (!own) return;
      const r = el.getBoundingClientRect(); const c = getComputedStyle(el);
      if (r.width < 2 || r.height < 2) return;
      let txt = own;
      if (c.textTransform === 'uppercase') txt = txt.toUpperCase();
      slide.testi.push({x: r.left - sr.left, y: r.top - sr.top, w: r.width, h: r.height, t: txt,
        ff: c.fontFamily, fs: parseFloat(c.fontSize), fw: c.fontWeight, fi: c.fontStyle, col: c.color,
        al: c.textAlign, lh: c.lineHeight, bg: c.backgroundColor,
        bordo: c.borderTopStyle !== 'none' && parseFloat(c.borderTopWidth) > 0 ? {st: c.borderTopStyle, col: c.borderTopColor} : null});
      if (el.tagName === 'LI') {
        const b = getComputedStyle(el, '::before');
        if (b.content && b.content !== 'none' && b.content !== 'normal' && !/counter/.test(b.content)) {
          slide.linee.push({x: r.left - sr.left, y: r.top - sr.top + parseFloat(c.fontSize) * 0.64, w: 14, h: 1.5, col: '#C48A3A'});
        }
        if (/counter/.test(b.content)) {
          const n = [...el.parentElement.children].indexOf(el) + 1;
          slide.testi.push({x: r.left - sr.left, y: r.top - sr.top, w: 50, h: 30, t: String(n).padStart(2, '0'),
            ff: 'Outfit', fs: 20, fw: '600', fi: 'normal', col: 'rgb(122, 79, 94)', al: 'left', lh: 'normal', bg: 'rgba(0, 0, 0, 0)', bordo: null});
        }
      }
    });
    sec.querySelectorAll('.filetto').forEach(el => { const r = el.getBoundingClientRect(); slide.linee.push({x: r.left - sr.left, y: r.top - sr.top, w: Math.max(1, r.width), h: r.height, col: '#C48A3A'}); });
    sec.querySelectorAll('.punto-fine').forEach(el => { const r = el.getBoundingClientRect(); slide.punti.push({x: r.left - sr.left, y: r.top - sr.top, w: r.width, h: r.height, col: '#FEBD59'}); });
    sec.querySelectorAll('.flusso li').forEach(el => { const r = el.getBoundingClientRect(); slide.linee.push({x: r.left - sr.left, y: r.top - sr.top, w: r.width, h: 3, col: '#C48A3A'}); });
    out.push(slide);
  });
  return out;
}
"""


def rgb(css: str) -> RGBColor | None:
    m = re.match(r"rgba?\((\d+),\s*(\d+),\s*(\d+)(?:,\s*([\d.]+))?\)", css or "")
    if not m:
        return None
    if m.group(4) is not None and float(m.group(4)) == 0:
        return None
    return RGBColor(int(m.group(1)), int(m.group(2)), int(m.group(3)))


def hexrgb(h: str) -> RGBColor:
    h = h.lstrip("#")
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def px(v: float) -> Emu:
    return Emu(int(round(v * EMU_PER_PX)))


def famiglia(ff: str) -> str:
    primo = ff.split(",")[0].strip().strip("'\"")
    return FONT.get(primo, primo)


def main() -> None:
    prs = Presentation()
    prs.slide_width = Emu(int(LARGH_IN * 914400))
    prs.slide_height = Emu(int(LARGH_IN * 914400 * H / W))
    vuoto = prs.slide_layouts[6]

    with sync_playwright() as p:
        br = p.chromium.launch()
        pg = br.new_page(viewport={"width": W, "height": H}, device_scale_factor=1)
        pg.goto(SORGENTE.as_uri())
        pg.wait_for_function("document.fonts.ready")
        pg.wait_for_timeout(500)
        misure = pg.evaluate(JS_MISURA, {"raster": ", ".join(RASTER), "testo": TESTO})
        sezioni = pg.query_selector_all("section.slide")

        for m, sec in zip(misure, sezioni):
            sl = prs.slides.add_slide(vuoto)
            fondo = rgb(m["bg"])
            if fondo is not None:
                sl.background.fill.solid()
                sl.background.fill.fore_color.rgb = fondo

            shot = Image.open(io.BytesIO(sec.screenshot()))

            for r in m["raster"]:
                box = (max(0, int(r["x"])), max(0, int(r["y"])), min(W, int(r["x"] + r["w"])), min(H, int(r["y"] + r["h"])))
                if box[2] - box[0] < 2 or box[3] - box[1] < 2:
                    continue
                crop = shot.crop(box)
                buf = io.BytesIO()
                if crop.width * crop.height > 300_000:
                    crop.convert("RGB").save(buf, "JPEG", quality=88)
                else:
                    crop.save(buf, "PNG", optimize=True)
                buf.seek(0)
                sl.shapes.add_picture(buf, px(box[0]), px(box[1]), px(box[2] - box[0]), px(box[3] - box[1]))

            for l in m["linee"]:
                sh = sl.shapes.add_shape(MSO_SHAPE.RECTANGLE, px(l["x"]), px(l["y"]), px(l["w"]), px(max(l["h"], 1)))
                sh.fill.solid(); sh.fill.fore_color.rgb = hexrgb(l["col"]); sh.line.fill.background()
            for d in m["punti"]:
                sh = sl.shapes.add_shape(MSO_SHAPE.OVAL, px(d["x"]), px(d["y"]), px(d["w"]), px(d["h"]))
                sh.fill.solid(); sh.fill.fore_color.rgb = hexrgb(d["col"]); sh.line.fill.background()

            for t in m["testi"]:
                tb = sl.shapes.add_textbox(px(t["x"] - 2), px(t["y"] - 2), px(t["w"] + 6), px(t["h"] + 6))
                tf = tb.text_frame
                tf.word_wrap = True
                tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
                tf.vertical_anchor = MSO_ANCHOR.TOP
                sfondo = rgb(t["bg"])
                if sfondo is not None:
                    tb.fill.solid(); tb.fill.fore_color.rgb = sfondo
                    tf.margin_left = tf.margin_right = px(12); tf.margin_top = tf.margin_bottom = px(6)
                if t["bordo"]:
                    tb.line.color.rgb = rgb(t["bordo"]["col"]) or hexrgb("#0999B3")
                    tb.line.width = Pt(1.5)
                    if t["bordo"]["st"] == "dashed":
                        tb.line.dash_style = MSO_LINE.DASH
                    tf.margin_left = tf.margin_right = px(18); tf.margin_top = tf.margin_bottom = px(8)
                for i, riga in enumerate(t["t"].split("\n")):
                    par = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                    par.alignment = {"center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT, "end": PP_ALIGN.RIGHT}.get(t["al"], PP_ALIGN.LEFT)
                    lh = t["lh"]
                    if lh and lh != "normal":
                        par.line_spacing = max(0.9, float(lh.replace("px", "")) / t["fs"])
                    run = par.add_run()
                    run.text = riga
                    f = run.font
                    f.name = famiglia(t["ff"])
                    f.size = Pt(t["fs"] * 0.75)
                    f.bold = int(t["fw"]) >= 600 if t["fw"].isdigit() else t["fw"] == "bold"
                    f.italic = t["fi"] == "italic"
                    col = rgb(t["col"])
                    if col is not None:
                        f.color.rgb = col

        br.close()

    USCITA.parent.mkdir(exist_ok=True)
    prs.save(USCITA)
    print(USCITA, round(USCITA.stat().st_size / 1e6, 2), "MB", len(prs.slides), "slide")


if __name__ == "__main__":
    main()
